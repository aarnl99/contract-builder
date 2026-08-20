"""
Core .docx manipulation engine.

Responsibilities:
  1. Render a .docx's content -- top-level paragraphs AND paragraphs inside
     table cells (recursing into nested tables), in document order -- into
     HTML where every run of text is wrapped in a
     <span data-p="{paragraph index}" data-r="{run index}"> so the browser
     can map a text selection back to an exact
     (container, paragraph, run, char-offset) location.
  2. Given such a selection, split/trim the underlying XML runs and replace
     the selected text with a "{{field_key}}" token, preserving the
     original run formatting (bold, italic, font, etc.) on either side of
     the token.
  3. Given a template's stored placeholders and a dict of user-supplied
     values, produce a finished .docx with every token replaced, wherever
     in the document it lives (body paragraph or table cell).

A "container" is either the document body itself or a single table cell.
Every container exposes `.paragraphs`, `.tables`, and `.iter_inner_content()`
(python-docx gives both Document and _Cell this interface), which is what
lets the same rendering/lookup code walk into cells and nested tables
without caring which kind of container it's in.

A "container path" is how we address a specific cell from the document
root: a list of [table_index, row_index, col_index] triples, one per level
of nesting. [] means "the top-level document body". table_index/row_index/
col_index are all local to their immediate parent container, matching
python-docx's own `.tables` / cell-grid indexing.

Known limitation: headers and footers are separate document parts and are
still not rendered or markable. Merged table cells are rendered as a single
cell without recomputing colspan/rowspan, so a merged region may look
slightly narrower/shorter in the marking view than in Word -- the content
itself is still fully present and markable.
"""
import copy
import html
import re
from typing import Dict, List, Optional

from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run as DocxRun

TOKEN_RE = re.compile(r"^\{\{([a-zA-Z0-9_]+)\}\}$")


def token_for(field_key: str) -> str:
    return "{{%s}}" % field_key


def load(path: str) -> Document:
    return Document(path)


def save(doc: Document, path: str) -> None:
    doc.save(path)


class MarkError(ValueError):
    pass


# ---------------------------------------------------------------------------
# Container helpers: walking body paragraphs + table cells uniformly
# ---------------------------------------------------------------------------

def _iter_unique_cells(table: Table):
    """Yield (row_idx, col_idx, cell) for each physically distinct cell in
    the table, addressed by the (row, col) of its first occurrence in
    reading order. A cell spanned by a horizontal/vertical merge is the
    same underlying XML element for every grid position it covers, so we
    skip it on repeat occurrences -- each physical cell is yielded once.

    Dedup key is the cell's underlying lxml element itself (`cell._tc`),
    NOT `id(cell._tc)`. python-docx hands back a fresh `_Cell` wrapper
    object on every `row.cells` access, so the wrapper (and the id() of
    whatever it points at) is only alive for that one iteration -- CPython
    is then free to reuse that same memory address for the *next* cell's
    wrapper, making two genuinely distinct cells collide on `id()` and get
    misidentified as a repeat of the same merged cell. That's not
    hypothetical: it reproduced on a plain, unmerged 2x2 table, silently
    dropping the last cell from rendering, marking, redlining, AND
    placeholder-value substitution at generation time (a real generated
    contract could go out with a literal unreplaced "{{field_key}}" token).
    Keying on the element itself avoids this: the set holds a real
    reference to each element for as long as iteration runs, so nothing
    it points at can be freed and its address reused underneath us."""
    seen = set()
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            key = cell._tc
            if key in seen:
                continue
            seen.add(key)
            yield r_idx, c_idx, cell


def _resolve_container(doc: Document, path: List[List[int]]):
    """Walk a container path from the document root and return the
    container (Document or _Cell) it points to."""
    container = doc
    for step in path:
        if len(step) != 3:
            raise MarkError("Invalid container path")
        t_idx, r_idx, c_idx = step
        tables = container.tables
        if t_idx < 0 or t_idx >= len(tables):
            raise MarkError(f"Invalid table index {t_idx}")
        table = tables[t_idx]
        found = None
        for rr, cc, cell in _iter_unique_cells(table):
            if rr == r_idx and cc == c_idx:
                found = cell
                break
        if found is None:
            raise MarkError(f"Invalid cell address ({r_idx}, {c_idx})")
        container = found
    return container


def _iter_all_paragraphs(container):
    """Recursively yield every paragraph in a container: its own top-level
    paragraphs plus, for every table it contains, every paragraph in every
    cell of that table (recursing into nested tables)."""
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            yield block
        elif isinstance(block, Table):
            for _, _, cell in _iter_unique_cells(block):
                yield from _iter_all_paragraphs(cell)


# ---------------------------------------------------------------------------
# Rendering: docx -> run-tagged HTML
# ---------------------------------------------------------------------------

def _path_str(path: List[List[int]]) -> str:
    return ";".join(",".join(str(x) for x in step) for step in path)


def _render_run_span(run, p_idx: int, r_idx: int, field_key: Optional[str] = None) -> str:
    text = run.text or ""
    escaped = html.escape(text, quote=False)
    style_bits = []
    if run.bold:
        style_bits.append("font-weight:bold")
    if run.italic:
        style_bits.append("font-style:italic")
    if run.underline:
        style_bits.append("text-decoration:underline")
    style_attr = f' style="{";".join(style_bits)}"' if style_bits else ""
    is_token = bool(TOKEN_RE.match(text.strip()))
    extra_class = ""
    extra_attr = ""
    if is_token:
        extra_class = " token"
    elif field_key:
        extra_class = " field-value"
        extra_attr = f' data-field-key="{html.escape(field_key, quote=True)}"'
    return f'<span class="run{extra_class}" data-p="{p_idx}" data-r="{r_idx}"{extra_attr}{style_attr}>{escaped}</span>'


def _render_paragraph_html(paragraph: Paragraph, p_idx: int, path: List[List[int]], field_lookup: Optional[Dict] = None) -> str:
    path_attr = f' data-path="{_path_str(path)}"' if path else ""
    runs = paragraph.runs
    if not runs:
        return f'<p class="para" data-p="{p_idx}"{path_attr}>&nbsp;</p>'
    path_key = _path_str(path)
    run_spans = [
        _render_run_span(run, p_idx, r_idx, field_lookup.get((path_key, p_idx, r_idx)) if field_lookup else None)
        for r_idx, run in enumerate(runs)
    ]
    return f'<p class="para" data-p="{p_idx}"{path_attr}>' + "".join(run_spans) + "</p>"


def _render_table_html(table: Table, t_idx: int, parent_path: List[List[int]], field_lookup: Optional[Dict] = None) -> str:
    rows: Dict[int, List[str]] = {}
    for r_idx, c_idx, cell in _iter_unique_cells(table):
        cell_path = parent_path + [[t_idx, r_idx, c_idx]]
        cell_html = _render_container_html(cell, cell_path, field_lookup)
        rows.setdefault(r_idx, []).append(cell_html or "&nbsp;")
    trs = []
    for r_idx in sorted(rows):
        tds = "".join(f'<td class="doc-td">{cell_html}</td>' for cell_html in rows[r_idx])
        trs.append(f"<tr>{tds}</tr>")
    return f'<table class="doc-table" data-table="{t_idx}"><tbody>' + "".join(trs) + "</tbody></table>"


def _render_container_html(container, path: List[List[int]], field_lookup: Optional[Dict] = None) -> str:
    parts = []
    p_idx = 0
    t_idx = 0
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            parts.append(_render_paragraph_html(block, p_idx, path, field_lookup))
            p_idx += 1
        elif isinstance(block, Table):
            parts.append(_render_table_html(block, t_idx, path, field_lookup))
            t_idx += 1
    return "\n".join(parts)


def render_paragraphs_html(doc: Document, field_positions: Optional[List[Dict]] = None) -> str:
    """Render the whole document body -- paragraphs and tables, in document
    order, recursing into table cells and nested tables -- as HTML with
    data-p/data-r spans for exact-offset selection mapping.

    If field_positions is given (a list of {"path", "p", "r", "field_key"}
    dicts, as produced by fill_template_tracked), the runs at those exact
    positions get a `field-value` class and `data-field-key` attribute so a
    client-facing page can find and highlight each drafted field inline,
    in place, instead of needing a separate form."""
    field_lookup = None
    if field_positions:
        field_lookup = {(fp["path"], fp["p"], fp["r"]): fp["field_key"] for fp in field_positions}
    return _render_container_html(doc, [], field_lookup)


# ---------------------------------------------------------------------------
# Marking: replace a selection with a {{token}}
# ---------------------------------------------------------------------------

def _clone_run_with_text(anchor_run, text: str):
    """Deep-copy anchor_run's XML (to keep its formatting) and set new text."""
    new_r = copy.deepcopy(anchor_run._r)
    new_run = DocxRun(new_r, anchor_run._parent)
    new_run.text = text
    return new_run


def _split_run_insert_token(run, start: int, end: int, token_text: str) -> None:
    """The FIRST (visually leftmost) segment of a selection: split this run
    into [before][TOKEN][after], preserving formatting on before/after."""
    text = run.text or ""
    before = text[:start]
    after = text[end:]
    r_element = run._r

    if before:
        run.text = before

    new_run_middle = _clone_run_with_text(run, token_text)
    r_element.addnext(new_run_middle._r)

    if after:
        new_run_after = _clone_run_with_text(run, after)
        new_run_middle._r.addnext(new_run_after._r)

    if not before:
        parent_el = r_element.getparent()
        parent_el.remove(r_element)


def _remove_range_from_run(run, start: int, end: int) -> None:
    """A subsequent (non-first) segment of a multi-run selection: just cut
    out the covered characters, no token inserted here."""
    text = run.text or ""
    new_text = text[:start] + text[end:]
    if new_text:
        run.text = new_text
    else:
        r_element = run._r
        parent_el = r_element.getparent()
        parent_el.remove(r_element)


def _utf16_offset_to_codepoint(text: str, utf16_offset: int) -> int:
    """Convert a UTF-16 code-unit offset (as produced by browser
    Range.startOffset/endOffset and Node.textContent.length -- see app.js's
    and share.js's computeCursorPosition/computeSelectionSegments) into the
    equivalent Python string index into `text`.

    Python strings index by codepoint; JavaScript strings index by UTF-16
    code unit. Every character is 1:1 EXCEPT characters outside the Basic
    Multilingual Plane (ord(ch) > 0xFFFF -- some emoji, some rare CJK
    Extension B+ ideographs, mathematical alphanumeric symbols, etc.), which
    take 2 UTF-16 code units but remain a single Python codepoint. A run
    containing even one such character before the split point makes a raw
    JS offset diverge from the correct Python index by one unit per such
    character -- silently slicing at the wrong boundary (or splitting a
    surrogate pair) instead of erroring. Bug tracker #49.
    """
    if utf16_offset <= 0:
        return 0
    units = 0
    for i, ch in enumerate(text):
        units += 2 if ord(ch) > 0xFFFF else 1
        if units >= utf16_offset:
            return i + 1
    # utf16_offset exceeds the text's total UTF-16 length -- a genuinely
    # out-of-range offset (bad/stale client data), not a legitimate "select
    # to the end of the run" case (that's already handled above: the loop's
    # final iteration returns len(text) once cumulative units reach
    # utf16_offset exactly). Returning it UNCHANGED here, rather than
    # clamping to len(text), matters: a codepoint contributes at most as
    # many UTF-16 units as itself, so the text's total UTF-16 length is
    # always >= len(text) -- meaning utf16_offset here is guaranteed >
    # len(text) too, so the caller's existing bounds check
    # (0 <= start <= end <= len(text)) still correctly rejects it as
    # invalid instead of silently clamping a bad offset into a valid one.
    return utf16_offset


def _normalize_segment_offsets(seg: dict, text: str) -> dict:
    """Return a copy of `seg` with start/end converted from UTF-16 code-unit
    offsets (as sent by the browser) to Python codepoint offsets against the
    given run's CURRENT text -- see _utf16_offset_to_codepoint. Every
    consumer below indexes `run.text` by codepoint, so this must run before
    any bounds-check or slicing sees the raw offsets. Bug tracker #49."""
    norm = dict(seg)
    norm["start"] = _utf16_offset_to_codepoint(text, seg["start"])
    norm["end"] = _utf16_offset_to_codepoint(text, seg["end"])
    return norm


def extract_text_at(doc: Document, container_path: Optional[List[List[int]]], paragraph_index: int, segments: List[dict]) -> str:
    """Read back the exact text currently at a given (container, paragraph,
    run-segments) location. Used server-side to derive a redline edit's
    "original text" from the live document instead of trusting whatever the
    browser sent, and to reject a stale/invalid selection (the document
    changed shape since the client loaded it) as a MarkError up front rather
    than reading garbage or silently misapplying it later."""
    container = _resolve_container(doc, container_path or [])
    paragraphs = container.paragraphs
    if paragraph_index < 0 or paragraph_index >= len(paragraphs):
        raise MarkError("Invalid paragraph index")
    runs = paragraphs[paragraph_index].runs
    parts = []
    for seg in sorted(segments, key=lambda s: (s["r"], s["start"])):
        r_idx = seg["r"]
        if r_idx < 0 or r_idx >= len(runs):
            raise MarkError(f"Invalid run index {r_idx}")
        text = runs[r_idx].text or ""
        seg = _normalize_segment_offsets(seg, text)
        if not (0 <= seg["start"] <= seg["end"] <= len(text)):
            raise MarkError(f"Invalid offsets for run {r_idx}")
        parts.append(text[seg["start"]:seg["end"]])
    return "".join(parts)


def _apply_paragraph_group(paragraph: Paragraph, group: List[dict]) -> None:
    """Rebuild one paragraph's runs, applying every edit in `group` (each
    {"segments": [...], "new_text": str}) in a single pass over the
    paragraph's ORIGINAL run list. All edits' segments were captured by the
    browser against that same original layout, so this must not re-query
    paragraph.runs between edits -- doing so one at a time would shift run
    indices out from under any edit sharing this paragraph."""
    original_runs = list(paragraph.runs)

    tagged = []  # (run_index, start, end, group_index)
    for g_idx, e in enumerate(group):
        for seg in e["segments"]:
            r_idx = seg["r"]
            if r_idx < 0 or r_idx >= len(original_runs):
                raise MarkError(f"Invalid run index {r_idx}")
            run_text = original_runs[r_idx].text or ""
            seg = _normalize_segment_offsets(seg, run_text)
            run_len = len(run_text)
            if not (0 <= seg["start"] <= seg["end"] <= run_len):
                raise MarkError(f"Invalid offsets for run {r_idx}")
            tagged.append((r_idx, seg["start"], seg["end"], g_idx))
    tagged.sort(key=lambda x: (x[0], x[1]))

    for i in range(1, len(tagged)):
        pr, _, pe, _ = tagged[i - 1]
        r, s, _, _ = tagged[i]
        if r == pr and s < pe:
            raise MarkError("Overlapping redline selections can't be applied together")

    segs_by_run: Dict[int, List[tuple]] = {}
    for r_idx, s, e, g_idx in tagged:
        segs_by_run.setdefault(r_idx, []).append((s, e, g_idx))

    emitted = set()  # group indices whose new_text has already been written once
    for r_idx, run in enumerate(original_runs):
        segs = segs_by_run.get(r_idx)
        if not segs:
            continue
        segs.sort(key=lambda x: x[0])
        text = run.text or ""
        pieces = []
        cursor = 0
        for s, e, g_idx in segs:
            if s > cursor:
                pieces.append(text[cursor:s])
            if g_idx not in emitted:
                pieces.append(group[g_idx]["new_text"])
                emitted.add(g_idx)
            cursor = e
        if cursor < len(text):
            pieces.append(text[cursor:])

        r_element = run._r
        first = pieces[0] if pieces else ""
        run.text = first
        anchor = r_element
        for piece in pieces[1:]:
            if not piece:
                continue
            new_run = _clone_run_with_text(run, piece)
            anchor.addnext(new_run._r)
            anchor = new_run._r
        if not first:
            parent_el = r_element.getparent()
            parent_el.remove(r_element)


def apply_text_edits(doc: Document, edits: List[dict]) -> None:
    """Apply a batch of redline text replacements directly onto an
    already-generated document -- unlike mark_placeholder, which inserts a
    {{token}} into a *template*, this writes literal replacement text in
    place. Each edit is {"container_path", "paragraph_index",
    "segments": [{"r","start","end"}, ...], "new_text"}. Edits are grouped
    by paragraph so multiple edits landing in the same paragraph are applied
    together in one pass (see _apply_paragraph_group) instead of
    invalidating each other's recorded run indices."""
    groups: Dict[tuple, List[dict]] = {}
    path_by_key: Dict[tuple, List[List[int]]] = {}
    for e in edits:
        cp = e.get("container_path") or []
        key = (_path_str(cp), e["paragraph_index"])
        path_by_key[key] = cp
        groups.setdefault(key, []).append(e)

    for key, group in groups.items():
        container_path = path_by_key[key]
        container = _resolve_container(doc, container_path)
        paragraphs = container.paragraphs
        p_idx = key[1]
        if p_idx < 0 or p_idx >= len(paragraphs):
            raise MarkError("Invalid paragraph index")
        _apply_paragraph_group(paragraphs[p_idx], group)


def mark_placeholder(
    doc: Document,
    paragraph_index: int,
    segments: List[dict],
    field_key: str,
    container_path: Optional[List[List[int]]] = None,
) -> None:
    """
    segments: list of {"r": run_index, "start": int, "end": int}, all within
    the same paragraph, referring to ORIGINAL (pre-edit) run indices/offsets
    as captured by the browser at selection time.

    container_path: [] (or None) for a top-level body paragraph, or a list
    of [table_index, row_index, col_index] triples locating the table cell
    the paragraph lives in (see module docstring).
    """
    container = _resolve_container(doc, container_path or [])
    paragraphs = container.paragraphs
    if paragraph_index < 0 or paragraph_index >= len(paragraphs):
        raise MarkError("Invalid paragraph index")
    paragraph = paragraphs[paragraph_index]
    original_runs = list(paragraph.runs)

    if not segments:
        raise MarkError("No selection segments provided")

    segments = sorted(segments, key=lambda s: s["r"])

    normalized_segments = []
    for seg in segments:
        r_idx = seg["r"]
        if r_idx < 0 or r_idx >= len(original_runs):
            raise MarkError(f"Invalid run index {r_idx}")
        run_text = original_runs[r_idx].text or ""
        seg = _normalize_segment_offsets(seg, run_text)
        if not (0 <= seg["start"] <= seg["end"] <= len(run_text)):
            raise MarkError(f"Invalid offsets for run {r_idx}")
        normalized_segments.append(seg)
    segments = normalized_segments

    token_text = token_for(field_key)
    first_seg = segments[0]
    _split_run_insert_token(
        original_runs[first_seg["r"]], first_seg["start"], first_seg["end"], token_text
    )
    for seg in segments[1:]:
        _remove_range_from_run(original_runs[seg["r"]], seg["start"], seg["end"])


# ---------------------------------------------------------------------------
# Generation: fill tokens with values
# ---------------------------------------------------------------------------

def fill_template(doc: Document, values: Dict[str, str]) -> int:
    """Replace every {{field_key}} run with its value, wherever it lives in
    the document (body paragraph or table cell, at any nesting depth).
    Returns count filled."""
    filled = 0
    for paragraph in _iter_all_paragraphs(doc):
        for run in paragraph.runs:
            text = run.text or ""
            m = TOKEN_RE.match(text.strip())
            if m:
                key = m.group(1)
                if key in values:
                    run.text = values[key]
                    filled += 1
    return filled


def fill_template_tracked(doc: Document, values: Dict[str, str]) -> List[Dict]:
    """Same fill as fill_template, but walks the document with the exact
    same container/path/paragraph-index/run-index scheme as
    render_paragraphs_html, and records where each field ended up. The
    returned list -- [{"path", "p", "r", "field_key"}, ...] -- can be
    stored alongside a generated contract and handed back into
    render_paragraphs_html later so a rendered copy of that same doc (e.g.
    on the client-facing share page) can highlight each field inline,
    without needing to content-match filled-in text back to a field."""
    positions: List[Dict] = []

    def walk(container, path: List[List[int]]) -> None:
        path_key = _path_str(path)
        p_idx = 0
        t_idx = 0
        for block in container.iter_inner_content():
            if isinstance(block, Paragraph):
                for r_idx, run in enumerate(block.runs):
                    text = run.text or ""
                    m = TOKEN_RE.match(text.strip())
                    if m:
                        key = m.group(1)
                        if key in values:
                            run.text = values[key]
                            positions.append({"path": path_key, "p": p_idx, "r": r_idx, "field_key": key})
                p_idx += 1
            elif isinstance(block, Table):
                for r_idx, c_idx, cell in _iter_unique_cells(block):
                    walk(cell, path + [[t_idx, r_idx, c_idx]])
                t_idx += 1

    walk(doc, [])
    return positions


def list_tokens(doc: Document) -> List[str]:
    """Utility: scan the whole doc (including table cells) and return all
    {{field_key}} tokens found."""
    found = []
    for paragraph in _iter_all_paragraphs(doc):
        for run in paragraph.runs:
            m = TOKEN_RE.match((run.text or "").strip())
            if m:
                found.append(m.group(1))
    return found
