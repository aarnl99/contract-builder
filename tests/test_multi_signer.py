"""Tests for sending a document to any number of signers (CEO, CFO, client...),
with custom role labels and an optional sign-in-order mode.

Covers: three signers end to end, parallel vs ordered emailing (the next
signer is only emailed when the previous one signs), the "not your turn yet"
gate on the public signing page, role/email validation, and the cross-checks
between placed fields and the signer list. The original client/sender body
shape must keep working.

Run with:  venv/bin/pytest tests/
"""
import hashlib
import hmac
import json
import os
import sys
import tempfile
import time

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

if "app.db" not in sys.modules:
    os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="rotely_test_multi_signer_")

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlmodel import Session, select  # noqa: E402

from app import main as m  # noqa: E402
from app import docuseal_engine as ds  # noqa: E402
from app.db import engine, init_db  # noqa: E402
from app.auth import hash_password  # noqa: E402
from app.models import GeneratedContract, SigningRequest, SigningRequestSigner, User  # noqa: E402

init_db()

SECRET = "whsec_test_multi_signer"

import itertools  # noqa: E402

_id_counter = itertools.count(700000, 100)

SIGNERS3 = [
    {"name": "Casey CEO", "email": "ceo@x.test", "role": "CEO"},
    {"name": "Fran CFO", "email": "cfo@x.test", "role": "CFO"},
    {"name": "Jamie Client", "email": "client@x.test", "role": "Client"},
]


@pytest.fixture()
def env(monkeypatch):
    """A logged-in owner, a document with three '______' lines, a fake
    DocuSeal that records what it was asked to create, and captured email."""
    monkeypatch.setattr(ds, "DOCUSEAL_WEBHOOK_SECRET", SECRET)
    monkeypatch.setattr(m.ds, "get_submission_documents", lambda submission_id: {"documents": []})
    monkeypatch.setattr(m, "_notify", lambda *a, **k: None)

    sent = []
    monkeypatch.setattr(m.ee, "send_email", lambda to, subject, body, *a, **k: sent.append(to))

    created = []

    def fake_create(*, name, file_bytes, file_name, submitters, variables=None, order="preserved"):
        created.append({"submitters": submitters, "order": order})
        # DocuSeal ids are globally unique; keep them so across tests too,
        # since all tests share one database.
        base = next(_id_counter)
        return {
            "id": base,
            "submitters": [
                {"id": base + i + 1, "email": s["email"], "embed_src": "https://docuseal.test/s/%d" % (base + i + 1)}
                for i, s in enumerate(submitters)
            ],
        }

    monkeypatch.setattr(m.ds, "create_submission_from_docx", fake_create)

    doc = Document()
    for label in ("CEO", "CFO", "Client"):
        p = doc.add_paragraph()
        p.add_run(label + " signs: ")
        p.add_run("______")
    path = os.path.join(tempfile.mkdtemp(prefix="rotely_ms_doc_"), "contract.docx")
    doc.save(path)

    with Session(engine) as s:
        u = User(email="ms-%d@x.test" % time.time_ns(), password_hash=hash_password("pw"),
                 email_verified=True, name="Owner")
        s.add(u)
        s.commit()
        s.refresh(u)
        gc = GeneratedContract(owner_id=u.id, template_name="t", name="Big Contract",
                               file_path=path, values_json="[]")
        s.add(gc)
        s.commit()
        s.refresh(gc)
        email, gid = u.email, gc.id

    c = TestClient(m.app)
    assert c.post("/api/login", json={"email": email, "password": "pw"}).status_code == 200
    c.gid = gid
    c.sent = sent
    c.created = created
    return c


def place(c, paragraph_index, role, field_type="signature"):
    return c.post(f"/api/generated/{c.gid}/signature-fields", json={
        "paragraph_index": paragraph_index,
        "segments": [{"r": 1, "start": 0, "end": 6}],
        "expected_text": "______",
        "role": role,
        "field_type": field_type,
    })


def place_all(c):
    for i, s in enumerate(SIGNERS3):
        assert place(c, i, s["role"]).status_code == 200


def send(c, **body):
    return c.post(f"/api/generated/{c.gid}/signature-request", json=body)


def webhook(c, event, submitter_id):
    body = json.dumps({"event_type": event, "data": {"id": submitter_id}}, separators=(",", ":")).encode()
    ts = str(int(time.time()))
    sig = hmac.new(SECRET.encode(), (ts + ".").encode() + body, hashlib.sha256).hexdigest()
    return c.post("/api/webhooks/docuseal", content=body,
                  headers={"Content-Type": "application/json", "X-Docuseal-Signature": ts + "." + sig})


# ---- three signers, everyone at once -------------------------------------

def test_three_signers_all_at_once(env):
    place_all(env)
    r = send(env, signers=SIGNERS3)
    assert r.status_code == 200, r.text
    data = r.json()
    assert [s["role"] for s in data["signers"]] == ["CEO", "CFO", "Client"]
    assert [s["position"] for s in data["signers"]] == [0, 1, 2]
    assert data["sign_in_order"] is False
    assert all(s["your_turn"] for s in data["signers"])

    assert [x["role"] for x in env.created[0]["submitters"]] == ["CEO", "CFO", "Client"]
    assert env.created[0]["order"] == "random"
    assert sorted(env.sent) == sorted(s["email"] for s in SIGNERS3)


# ---- sign in order --------------------------------------------------------

def test_sign_in_order_emails_one_at_a_time(env):
    place_all(env)
    r = send(env, signers=SIGNERS3, sign_in_order=True)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["sign_in_order"] is True
    assert env.created[0]["order"] == "preserved"
    # only the first signer is emailed up front
    assert env.sent == ["ceo@x.test"]
    assert [s["your_turn"] for s in data["signers"]] == [True, False, False]

    with Session(engine) as s:
        rows = s.exec(select(SigningRequestSigner).where(
            SigningRequestSigner.signing_request_id == data["id"])).all()
        by_role = {x.role: x for x in rows}

    # first signer signs -> the CFO (and only the CFO) is emailed next
    assert webhook(env, "form.completed", by_role["CEO"].docuseal_submitter_id).status_code == 200
    assert env.sent == ["ceo@x.test", "cfo@x.test"]

    # a duplicate webhook must not email the CFO again
    assert webhook(env, "form.completed", by_role["CEO"].docuseal_submitter_id).status_code == 200
    assert env.sent == ["ceo@x.test", "cfo@x.test"]

    assert webhook(env, "form.completed", by_role["CFO"].docuseal_submitter_id).status_code == 200
    assert env.sent == ["ceo@x.test", "cfo@x.test", "client@x.test"]


def test_later_signer_cannot_sign_before_their_turn(env):
    place_all(env)
    data = send(env, signers=SIGNERS3, sign_in_order=True).json()
    with Session(engine) as s:
        client_signer = s.exec(select(SigningRequestSigner).where(
            SigningRequestSigner.signing_request_id == data["id"],
            SigningRequestSigner.role == "Client")).first()
        token = client_signer.token

    pub = TestClient(m.app)  # public page: no login
    status = pub.get(f"/api/sign/{token}").json()
    assert status["your_turn"] is False
    assert status["waiting_on"] == ["Casey CEO", "Fran CFO"]
    assert status["embed_src"] is None

    r = pub.post(f"/api/sign/{token}/consent", json={"consent": True, "photo": "data:image/png;base64,AAAA"})
    assert r.status_code == 400
    assert "turn" in r.json()["detail"]


def test_single_signer_ignores_order_flag(env):
    place(env, 0, "CEO")
    r = send(env, signers=[SIGNERS3[0]], sign_in_order=True)
    assert r.status_code == 200, r.text
    assert r.json()["sign_in_order"] is False
    assert env.created[0]["order"] == "random"


# ---- validation -----------------------------------------------------------

def test_duplicate_email_rejected_case_insensitive(env):
    dup = [SIGNERS3[0], {**SIGNERS3[1], "email": "CEO@x.test"}]
    r = send(env, signers=dup)
    assert r.status_code == 400
    assert "twice" in r.json()["detail"]
    assert not env.created


def test_duplicate_role_rejected_case_insensitive(env):
    dup = [SIGNERS3[0], {**SIGNERS3[1], "role": "ceo"}]
    r = send(env, signers=dup)
    assert r.status_code == 400
    assert "role" in r.json()["detail"].lower()


@pytest.mark.parametrize("bad_role", ["CEO;role=x", "{{CEO}}", "", "  ", "a" * 41, "CEO=1"])
def test_unsafe_or_empty_role_rejected(env, bad_role):
    r = send(env, signers=[{**SIGNERS3[0], "role": bad_role}])
    assert r.status_code == 400
    assert not env.created


def test_role_whitespace_is_normalised(env):
    place(env, 0, "Co  signer")  # collapses to "Co signer"
    r = send(env, signers=[{**SIGNERS3[0], "role": "  Co   signer "}])
    assert r.status_code == 200, r.text
    assert r.json()["signers"][0]["role"] == "Co signer"


def test_too_many_signers_rejected(env):
    many = [{"name": "N%d" % i, "email": "n%d@x.test" % i, "role": "R%d" % i} for i in range(11)]
    r = send(env, signers=many)
    assert r.status_code == 400
    assert "10" in r.json()["detail"]


def test_bad_email_rejected(env):
    r = send(env, signers=[{**SIGNERS3[0], "email": "not-an-email"}])
    assert r.status_code == 400


def test_place_field_rejects_unsafe_role(env):
    assert place(env, 0, "CEO;role=x").status_code == 400


# ---- fields <-> signer list ----------------------------------------------

def test_field_for_signer_not_on_list_blocks_send(env):
    place_all(env)
    r = send(env, signers=SIGNERS3[:2])  # Client dropped, but has a field
    assert r.status_code == 400
    assert "Client (1)" in r.json()["detail"]
    assert not env.created


def test_signer_without_a_field_blocks_send(env):
    place(env, 0, "CEO")
    place(env, 1, "CFO")
    r = send(env, signers=SIGNERS3)  # Client has no field
    assert r.status_code == 400
    assert "Client" in r.json()["detail"]
    assert not env.created


# ---- backwards compatibility ---------------------------------------------

def test_legacy_client_sender_body_still_works(env):
    r = send(env, client_name="Jamie Client", client_email="client@x.test",
             include_sender=True, sender_name="Owner", sender_email="owner@x.test")
    assert r.status_code == 200, r.text
    assert [s["role"] for s in r.json()["signers"]] == ["Client", "Sender"]
    assert [s["position"] for s in r.json()["signers"]] == [0, 1]
    assert sorted(env.sent) == ["client@x.test", "owner@x.test"]


def test_legacy_body_requires_client(env):
    assert send(env).status_code == 400


# ---- cancel, then start over ---------------------------------------------

def test_can_send_again_after_cancel_with_different_people(env, monkeypatch):
    monkeypatch.setattr(m.ds, "archive_submission", lambda submission_id: {})
    place_all(env)
    first = send(env, signers=SIGNERS3)
    assert first.status_code == 200, first.text
    # can't double-send while one is live
    assert send(env, signers=SIGNERS3).status_code == 400

    assert env.post(f"/api/generated/{env.gid}/signature-request/cancel").status_code == 200
    latest = env.get(f"/api/generated/{env.gid}/signature-request").json()
    assert latest["status"] == "cancelled"

    # same roles (their fields are kept), a different CFO
    again = [SIGNERS3[0], {"name": "New CFO", "email": "newcfo@x.test", "role": "CFO"}, SIGNERS3[2]]
    second = send(env, signers=again, sign_in_order=True)
    assert second.status_code == 200, second.text
    assert second.json()["id"] != first.json()["id"]
    latest = env.get(f"/api/generated/{env.gid}/signature-request").json()
    assert latest["id"] == second.json()["id"] and latest["status"] == "pending"
    assert [s["email"] for s in latest["signers"]] == ["ceo@x.test", "newcfo@x.test", "client@x.test"]


def test_old_cancelled_links_stay_dead_after_resend(env, monkeypatch):
    monkeypatch.setattr(m.ds, "archive_submission", lambda submission_id: {})
    place_all(env)
    first = send(env, signers=SIGNERS3).json()
    with Session(engine) as s:
        old_token = s.exec(select(SigningRequestSigner).where(
            SigningRequestSigner.signing_request_id == first["id"],
            SigningRequestSigner.role == "CFO")).first().token
    env.post(f"/api/generated/{env.gid}/signature-request/cancel")
    send(env, signers=SIGNERS3)
    pub = TestClient(m.app)
    st = pub.get(f"/api/sign/{old_token}").json()
    assert st["request_status"] == "cancelled"
    r = pub.post(f"/api/sign/{old_token}/consent", json={"consent": True, "photo": "data:image/png;base64,AAAA"})
    assert r.status_code == 400
