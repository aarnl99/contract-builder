"""Unit tests for the highlight-to-field signature flow (docx_engine side).

Covers apply_signature_field_tags: the function that splices DocuSeal
text tags into a throwaway copy of the document in place of the exact
ranges the owner highlighted.

Run with:  venv/bin/pytest tests/
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from docx import Document
from app import docx_engine as de


def make_sig_doc():
    """Paragraph 0: three runs. Paragraph 1: plain. Table with one cell."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Please sign: ")
    p.add_run("______")
    p.add_run(" and date: ______.")
    doc.add_paragraph("Thank you for your business.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell sig: ______ done"
    return doc


def field(paragraph_index=0, table_path="", segments=None, role="Client",
          field_type="signature", label="Client Signature 1", original_text="______"):
    return {
        "paragraph_index": paragraph_index,
        "table_path": table_path,
        "segments": segments if segments is not None else [{"r": 1, "start": 0, "end": 6}],
        "role": role,
        "field_type": field_type,
        "label": label,
        "original_text": original_text,
    }


def full_text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_single_run_tag_replaces_selection():
    doc = make_sig_doc()
    de.apply_signature_field_tags(doc, [field()])
    text = full_text(doc)
    assert "{{Client Signature 1;type=signature;role=Client;required=true}}" in text
    # the ______ is gone, surrounding text intact
    assert "Please sign: " in text
    assert " and date: ______." in text
    assert text.count("______") == 1  # only the date one remains


def test_tag_format_for_each_type_and_role():
    assert de.signature_field_tag("Client Signature 1", "signature", "Client") == \
        "{{Client Signature 1;type=signature;role=Client;required=true}}"
    assert de.signature_field_tag("Sender Initials 1", "initials", "Sender") == \
        "{{Sender Initials 1;type=initials;role=Sender;required=true}}"
    assert de.signature_field_tag("Client Date 1", "date", "Client") == \
        "{{Client Date 1;type=date;role=Client;required=true}}"
    assert de.signature_field_tag("Client Text 1", "text", "Client") == \
        "{{Client Text 1;type=text;role=Client;required=true}}"


def test_multi_run_selection_tag_at_first_segment_rest_removed():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("AB")
    p.add_run("CD")
    p.add_run("EF")
    # highlight "BCDE": run0[1,2) + run1[0,2) + run2[0,1)
    f = field(segments=[
        {"r": 0, "start": 1, "end": 2},
        {"r": 1, "start": 0, "end": 2},
        {"r": 2, "start": 0, "end": 1},
    ], original_text="BCDE", label="Client Signature 1")
    de.apply_signature_field_tags(doc, [f])
    text = doc.paragraphs[0].text
    assert text == "A{{Client Signature 1;type=signature;role=Client;required=true}}F"


def test_two_fields_same_run_no_overlap():
    doc = make_sig_doc()
    # run 2 is " and date: ______." -- the ______ starts at offset 11
    f1 = field()  # run 1 "______"
    f2 = field(
        segments=[{"r": 2, "start": 11, "end": 17}],
        label="Client Date 1", field_type="date", original_text="______",
    )
    de.apply_signature_field_tags(doc, [f1, f2])
    text = doc.paragraphs[0].text
    assert "{{Client Signature 1;type=signature;role=Client;required=true}}" in text
    assert "{{Client Date 1;type=date;role=Client;required=true}}" in text
    assert "______" not in text
    assert text.startswith("Please sign: ")
    assert " and date: " in text
    assert text.endswith(".")


def test_multi_run_field_plus_later_field_same_run_earlier_offset():
    # The regression case for per-run descending-offset application: field X
    # spans run0->run1, field Y sits in run1 at offsets BEFORE X's run1
    # segment. Both must land exactly.
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("0123456789")   # run 0
    p.add_run("ABCDEFGHIJ")   # run 1
    fX = field(segments=[
        {"r": 0, "start": 8, "end": 10},   # "89"
        {"r": 1, "start": 6, "end": 8},    # "GH"
    ], original_text="89GH", label="Client Signature 1")
    fY = field(segments=[
        {"r": 1, "start": 0, "end": 3},     # "ABC"
    ], original_text="ABC", label="Client Initials 1", field_type="initials")
    de.apply_signature_field_tags(doc, [fX, fY])
    text = doc.paragraphs[0].text
    assert text == (
        "01234567"
        "{{Client Signature 1;type=signature;role=Client;required=true}}"
        "{{Client Initials 1;type=initials;role=Client;required=true}}"
        "DEFIJ"
    )


def test_field_at_run_start_detaches_cleanly_with_other_field_same_run():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("______ and ______")
    f1 = field(segments=[{"r": 0, "start": 0, "end": 6}],
               label="Client Signature 1")
    f2 = field(segments=[{"r": 0, "start": 11, "end": 17}],
               label="Client Signature 2")
    de.apply_signature_field_tags(doc, [f1, f2])
    text = doc.paragraphs[0].text
    assert text == (
        "{{Client Signature 1;type=signature;role=Client;required=true}}"
        " and "
        "{{Client Signature 2;type=signature;role=Client;required=true}}"
    )


def test_table_cell_selection():
    doc = make_sig_doc()
    # cell text "Cell sig: ______ done" -- single run, ______ at [10,16)
    f = field(paragraph_index=0, table_path="0,0,0",
              segments=[{"r": 0, "start": 10, "end": 16}])
    de.apply_signature_field_tags(doc, [f])
    cell_text = doc.tables[0].cell(0, 0).text
    assert "{{Client Signature 1;type=signature;role=Client;required=true}}" in cell_text
    assert cell_text == ("Cell sig: "
                         "{{Client Signature 1;type=signature;role=Client;required=true}}"
                         " done")


def test_stale_selection_rejected():
    doc = make_sig_doc()
    # claim the ______ but the doc actually changed
    doc.paragraphs[0].runs[1].text = "XXXXXX"
    with pytest.raises(de.MarkError, match="no longer lines up"):
        de.apply_signature_field_tags(doc, [field()])


def test_invalid_offsets_rejected():
    doc = make_sig_doc()
    f = field(segments=[{"r": 1, "start": 0, "end": 999}])
    with pytest.raises(de.MarkError):
        de.apply_signature_field_tags(doc, [f])


def test_emoji_offsets_utf16_safe():
    # non-BMP chars: browser sends UTF-16 offsets, _normalize_segment_offsets
    # converts. "A\U0001F600B" -- emoji is 2 UTF-16 units.
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("A\U0001F600B______")
    # select "B______" : UTF-16 offsets: A=1 unit, emoji=2 units -> B at 3
    f = field(segments=[{"r": 0, "start": 3, "end": 10}],
              original_text="B______")
    de.apply_signature_field_tags(doc, [f])
    text = doc.paragraphs[0].text
    assert text == ("A\U0001F600"
                    "{{Client Signature 1;type=signature;role=Client;required=true}}")


def test_empty_fields_is_noop():
    doc = make_sig_doc()
    before = full_text(doc)
    de.apply_signature_field_tags(doc, [])
    assert full_text(doc) == before


def test_parse_container_path_str():
    assert de._parse_container_path_str("") == []
    assert de._parse_container_path_str("0,0,0") == [[0, 0, 0]]
    assert de._parse_container_path_str("0,0,0;1,2,3") == [[0, 0, 0], [1, 2, 3]]
    with pytest.raises(de.MarkError):
        de._parse_container_path_str("a,b,c")
