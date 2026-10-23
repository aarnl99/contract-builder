"""Tests for the per-contract photo check: the sender chooses none / selfie /
selfie + photo ID, the signer's page enforces it, photos are stored encrypted
and only the owner can view them, and DocuSeal's own record (IP, device,
times) is saved next to them.

Run with:  venv/bin/pytest tests/
"""
import base64
import hashlib
import hmac
import itertools
import json
import os
import sys
import tempfile
import time

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

if "app.db" not in sys.modules:
    os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="rotely_test_photo_check_")

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlmodel import Session, select  # noqa: E402

from app import main as m  # noqa: E402
from app import docuseal_engine as ds  # noqa: E402
from app import photostore as ps  # noqa: E402
from app.db import engine, init_db  # noqa: E402
from app.auth import hash_password  # noqa: E402
from app.models import GeneratedContract, SigningRequest, SigningRequestSigner, User  # noqa: E402

init_db()

SECRET = "whsec_test_photo_check"
_ids = itertools.count(900000, 100)

# Smallest byte strings that look like a real JPEG / PNG to the sniffer.
JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 40 + b"SELFIE-BYTES"
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 40 + b"ID-BYTES"


def data_url(raw, mime="image/jpeg"):
    return "data:%s;base64,%s" % (mime, base64.b64encode(raw).decode())


@pytest.fixture()
def env(monkeypatch):
    monkeypatch.setattr(ds, "DOCUSEAL_WEBHOOK_SECRET", SECRET)
    monkeypatch.setattr(m.ds, "get_submission_documents", lambda submission_id: {"documents": []})
    monkeypatch.setattr(m, "_notify", lambda *a, **k: None)
    sent = []
    monkeypatch.setattr(m.ee, "send_email", lambda to, subject, body, *a, **k: sent.append((to, body)))

    def fake_create(*, name, file_bytes, file_name, submitters, variables=None, order="preserved"):
        base = next(_ids)
        return {
            "id": base,
            "submitters": [
                {"id": base + i + 1, "email": s["email"], "embed_src": "https://docuseal.test/s/%d" % (base + i + 1)}
                for i, s in enumerate(submitters)
            ],
        }

    monkeypatch.setattr(m.ds, "create_submission_from_docx", fake_create)

    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Client signs: ")
    p.add_run("______")
    path = os.path.join(tempfile.mkdtemp(prefix="rotely_pc_doc_"), "contract.docx")
    doc.save(path)

    def make_owner():
        with Session(engine) as s:
            u = User(email="pc-%d@x.test" % time.time_ns(), password_hash=hash_password("pw"),
                     email_verified=True, name="Owner")
            s.add(u)
            s.commit()
            s.refresh(u)
            gc = GeneratedContract(owner_id=u.id, template_name="t", name="Photo Contract",
                                   file_path=path, values_json="[]")
            s.add(gc)
            s.commit()
            s.refresh(gc)
            email, gid = u.email, gc.id
        c = TestClient(m.app)
        assert c.post("/api/login", json={"email": email, "password": "pw"}).status_code == 200
        c.gid = gid
        return c

    c = make_owner()
    c.sent = sent
    c.make_owner = make_owner
    place = c.post(f"/api/generated/{c.gid}/signature-fields", json={
        "paragraph_index": 0, "segments": [{"r": 1, "start": 0, "end": 6}],
        "expected_text": "______", "role": "Client", "field_type": "signature",
    })
    assert place.status_code == 200, place.text
    return c


SIGNER = [{"name": "Jamie Client", "email": "jamie@x.test", "role": "Client"}]


def send(c, **extra):
    return c.post(f"/api/generated/{c.gid}/signature-request", json={"signers": SIGNER, **extra})


def token_of(c):
    with Session(engine) as s:
        sr = s.exec(select(SigningRequest).where(SigningRequest.generated_contract_id == c.gid)
                    .order_by(SigningRequest.id.desc())).first()
        sg = s.exec(select(SigningRequestSigner).where(SigningRequestSigner.signing_request_id == sr.id)).first()
        return sg.token, sg.id, sg.docuseal_submitter_id


def webhook(event, data):
    body = json.dumps({"event_type": event, "data": data}, separators=(",", ":")).encode()
    ts = str(int(time.time()))
    sig = hmac.new(SECRET.encode(), (ts + ".").encode() + body, hashlib.sha256).hexdigest()
    return TestClient(m.app).post("/api/webhooks/docuseal", content=body,
                                  headers={"Content-Type": "application/json", "X-Docuseal-Signature": ts + "." + sig})


# ---- the sender's choice ---------------------------------------------------

def test_default_is_no_photo_check(env):
    r = send(env)
    assert r.status_code == 200, r.text
    assert r.json()["photo_check"] == "none"


@pytest.mark.parametrize("level", ["none", "selfie", "selfie_id"])
def test_level_is_saved_and_reported(env, level):
    r = send(env, photo_check=level)
    assert r.status_code == 200, r.text
    assert r.json()["photo_check"] == level
    assert env.get(f"/api/generated/{env.gid}/signature-request").json()["photo_check"] == level


def test_unknown_level_is_refused_before_anything_is_sent(env):
    r = send(env, photo_check="retina_scan")
    assert r.status_code == 400
    assert env.get(f"/api/generated/{env.gid}/signature-request").json() == {"exists": False}


def test_email_wording_matches_the_level(env):
    send(env, photo_check="selfie_id")
    body = env.sent[-1][1]
    assert "selfie and a photo of your ID" in body


def test_email_for_no_check_does_not_mention_photos(env):
    send(env)
    assert "selfie" not in env.sent[-1][1].lower()


# ---- the signer's page enforces it ----------------------------------------

def test_public_status_tells_the_page_what_to_collect(env):
    send(env, photo_check="selfie_id")
    token, _, _ = token_of(env)
    assert TestClient(m.app).get(f"/api/sign/{token}").json()["photo_check"] == "selfie_id"


def test_none_needs_only_agreement(env):
    send(env)
    token, _, _ = token_of(env)
    pub = TestClient(m.app)
    assert pub.post(f"/api/sign/{token}/consent", json={"consent": False}).status_code == 400
    r = pub.post(f"/api/sign/{token}/consent", json={"consent": True})
    assert r.status_code == 200
    assert r.json()["embed_src"].startswith("https://docuseal.test/")


def test_selfie_level_requires_a_selfie(env):
    send(env, photo_check="selfie")
    token, _, _ = token_of(env)
    pub = TestClient(m.app)
    assert pub.post(f"/api/sign/{token}/consent", json={"consent": True}).status_code == 400
    r = pub.post(f"/api/sign/{token}/consent", json={"consent": True, "photo": data_url(JPEG)})
    assert r.status_code == 200, r.text


def test_selfie_id_level_requires_both_photos(env):
    send(env, photo_check="selfie_id")
    token, sid, _ = token_of(env)
    pub = TestClient(m.app)
    only_selfie = pub.post(f"/api/sign/{token}/consent", json={"consent": True, "photo": data_url(JPEG)})
    assert only_selfie.status_code == 400
    assert "ID" in only_selfie.json()["detail"]
    # The failed attempt must not leave a half-saved selfie or move the signer on.
    with Session(engine) as s:
        sg = s.get(SigningRequestSigner, sid)
        assert sg.status == "awaiting_consent" and not sg.photo_path
    ok = pub.post(f"/api/sign/{token}/consent",
                  json={"consent": True, "photo": data_url(JPEG), "id_photo": data_url(PNG, "image/png")})
    assert ok.status_code == 200, ok.text


@pytest.mark.parametrize("bad", ["", "not a data url", "data:image/jpeg;base64,@@@@", data_url(b"%PDF-1.4 hello"),
                                 data_url(b"<script>alert(1)</script>")])
def test_junk_photos_are_refused(env, bad):
    send(env, photo_check="selfie")
    token, _, _ = token_of(env)
    r = TestClient(m.app).post(f"/api/sign/{token}/consent", json={"consent": True, "photo": bad})
    assert r.status_code == 400


def test_header_lies_do_not_matter_only_the_bytes(env):
    send(env, photo_check="selfie")
    token, _, _ = token_of(env)
    r = TestClient(m.app).post(f"/api/sign/{token}/consent",
                               json={"consent": True, "photo": data_url(b"%PDF-1.4 hello", "image/jpeg")})
    assert r.status_code == 400


def test_oversized_photo_is_refused(env, monkeypatch):
    monkeypatch.setattr(ps, "MAX_IMAGE_BYTES", 100)
    send(env, photo_check="selfie")
    token, _, _ = token_of(env)
    big = JPEG + b"\x00" * 500
    r = TestClient(m.app).post(f"/api/sign/{token}/consent", json={"consent": True, "photo": data_url(big)})
    assert r.status_code == 400 and "too large" in r.json()["detail"]


# ---- storage: encrypted on disk, readable only by the owner ---------------

def test_photos_are_encrypted_on_disk(env):
    send(env, photo_check="selfie_id")
    token, sid, _ = token_of(env)
    TestClient(m.app).post(f"/api/sign/{token}/consent",
                           json={"consent": True, "photo": data_url(JPEG), "id_photo": data_url(PNG, "image/png")})
    with Session(engine) as s:
        sg = s.get(SigningRequestSigner, sid)
        for path, plain in ((sg.photo_path, JPEG), (sg.id_photo_path, PNG)):
            assert path.endswith(".enc")
            stored = open(path, "rb").read()
            assert plain not in stored
            assert stored != plain
        assert sg.photo_captured_at and sg.id_photo_captured_at


def test_owner_can_view_both_photos(env):
    send(env, photo_check="selfie_id")
    token, sid, _ = token_of(env)
    TestClient(m.app).post(f"/api/sign/{token}/consent",
                           json={"consent": True, "photo": data_url(JPEG), "id_photo": data_url(PNG, "image/png")})
    base = f"/api/generated/{env.gid}/signature-request/signers/{sid}/photo"
    selfie = env.get(base + "/selfie")
    assert selfie.status_code == 200 and selfie.content == JPEG
    assert selfie.headers["content-type"] == "image/jpeg"
    assert "no-store" in selfie.headers["cache-control"]
    assert selfie.headers["x-content-type-options"] == "nosniff"
    idp = env.get(base + "/id")
    assert idp.status_code == 200 and idp.content == PNG
    assert idp.headers["content-type"] == "image/png"


def test_another_account_cannot_view_the_photos(env):
    send(env, photo_check="selfie_id")
    token, sid, _ = token_of(env)
    TestClient(m.app).post(f"/api/sign/{token}/consent",
                           json={"consent": True, "photo": data_url(JPEG), "id_photo": data_url(PNG, "image/png")})
    other = env.make_owner()
    url = f"/api/generated/{env.gid}/signature-request/signers/{sid}/photo/selfie"
    assert other.get(url).status_code == 404
    # Pointing the other account's own document at this signer must not work either.
    assert other.get(f"/api/generated/{other.gid}/signature-request/signers/{sid}/photo/selfie").status_code == 404


def test_logged_out_visitor_cannot_view_the_photos(env):
    send(env, photo_check="selfie")
    token, sid, _ = token_of(env)
    TestClient(m.app).post(f"/api/sign/{token}/consent", json={"consent": True, "photo": data_url(JPEG)})
    anon = TestClient(m.app)
    r = anon.get(f"/api/generated/{env.gid}/signature-request/signers/{sid}/photo/selfie")
    assert r.status_code in (401, 403)


def test_bad_kind_or_missing_photo_is_404(env):
    send(env, photo_check="selfie")
    token, sid, _ = token_of(env)
    base = f"/api/generated/{env.gid}/signature-request/signers/{sid}/photo"
    assert env.get(base + "/selfie").status_code == 404  # nothing captured yet
    assert env.get(base + "/passport").status_code == 404
    assert env.get(base + "/id").status_code == 404


def test_legacy_plain_selfie_still_viewable(env):
    send(env, photo_check="selfie")
    token, sid, _ = token_of(env)
    legacy = os.path.join(m.UPLOADS_DIR, "legacy_%d.jpg" % sid)
    with open(legacy, "wb") as f:
        f.write(JPEG)
    with Session(engine) as s:
        sg = s.get(SigningRequestSigner, sid)
        sg.photo_path = legacy
        s.add(sg)
        s.commit()
    r = env.get(f"/api/generated/{env.gid}/signature-request/signers/{sid}/photo/selfie")
    assert r.status_code == 200 and r.content == JPEG


def test_photo_path_outside_uploads_is_never_served(env):
    send(env, photo_check="selfie")
    token, sid, _ = token_of(env)
    outside = os.path.join(tempfile.mkdtemp(prefix="rotely_outside_"), "x.jpg")
    with open(outside, "wb") as f:
        f.write(JPEG)
    with Session(engine) as s:
        sg = s.get(SigningRequestSigner, sid)
        sg.photo_path = outside
        s.add(sg)
        s.commit()
    assert env.get(f"/api/generated/{env.gid}/signature-request/signers/{sid}/photo/selfie").status_code == 404


def test_deleting_the_document_removes_the_photos(env):
    send(env, photo_check="selfie_id")
    token, sid, _ = token_of(env)
    TestClient(m.app).post(f"/api/sign/{token}/consent",
                           json={"consent": True, "photo": data_url(JPEG), "id_photo": data_url(PNG, "image/png")})
    with Session(engine) as s:
        sg = s.get(SigningRequestSigner, sid)
        paths = [sg.photo_path, sg.id_photo_path]
    assert all(os.path.exists(p) for p in paths)
    r = env.delete(f"/api/generated/{env.gid}")
    assert r.status_code == 200, r.text
    assert not any(os.path.exists(p) for p in paths)
    with Session(engine) as s:
        assert s.get(SigningRequestSigner, sid) is None


# ---- DocuSeal's own record beside ours ------------------------------------

def test_webhook_saves_docuseal_ip_device_and_times(env):
    send(env, photo_check="selfie")
    token, sid, submitter_id = token_of(env)
    TestClient(m.app).post(f"/api/sign/{token}/consent", json={"consent": True, "photo": data_url(JPEG)})
    r = webhook("form.completed", {
        "id": submitter_id,
        "ip": "203.0.113.9",
        "ua": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 Version/17.0 Safari/605.1.15",
        "opened_at": "2026-10-23T14:00:00.000Z",
        "completed_at": "2026-10-23T14:05:30.000Z",
    })
    assert r.status_code == 200
    payload = env.get(f"/api/generated/{env.gid}/signature-request").json()
    s0 = payload["signers"][0]
    assert s0["status"] == "completed"
    assert s0["ds_ip"] == "203.0.113.9"
    assert "Macintosh" in s0["ds_user_agent"]
    assert s0["ds_opened_at"].startswith("2026-10-23T14:00:00")
    assert s0["ds_completed_at"].startswith("2026-10-23T14:05:30")
    assert s0["has_photo"] is True and s0["has_id_photo"] is False


def test_webhook_without_details_still_completes(env):
    send(env)
    token, sid, submitter_id = token_of(env)
    assert webhook("form.completed", {"id": submitter_id}).status_code == 200
    s0 = env.get(f"/api/generated/{env.gid}/signature-request").json()["signers"][0]
    assert s0["status"] == "completed" and s0["ds_ip"] == "" and s0["ds_completed_at"] is None


def test_late_details_fill_in_an_already_completed_signer(env):
    send(env)
    token, sid, submitter_id = token_of(env)
    webhook("form.completed", {"id": submitter_id})
    webhook("form.completed", {"id": submitter_id, "ip": "198.51.100.7", "completed_at": "2026-10-23T09:00:00Z"})
    s0 = env.get(f"/api/generated/{env.gid}/signature-request").json()["signers"][0]
    assert s0["ds_ip"] == "198.51.100.7"


def test_refresh_path_also_saves_the_details(env, monkeypatch):
    send(env)
    token, sid, submitter_id = token_of(env)
    monkeypatch.setattr(m.ds, "get_submission", lambda sub_id: {
        "status": "pending",
        "submitters": [{"id": submitter_id, "status": "completed", "ip": "192.0.2.44",
                        "ua": "Mozilla/5.0 (Windows NT 10.0) Chrome/120.0 Safari/537.36",
                        "completed_at": "2026-10-23T10:00:00Z"}],
    })
    assert env.post(f"/api/generated/{env.gid}/signature-request/refresh").status_code == 200
    s0 = env.get(f"/api/generated/{env.gid}/signature-request").json()["signers"][0]
    assert s0["ds_ip"] == "192.0.2.44" and s0["ds_completed_at"].startswith("2026-10-23T10:00:00")


def test_garbage_details_are_ignored_safely(env):
    send(env)
    token, sid, submitter_id = token_of(env)
    webhook("form.completed", {"id": submitter_id, "ip": 12345, "ua": ["x"], "opened_at": "yesterday", "completed_at": 7})
    s0 = env.get(f"/api/generated/{env.gid}/signature-request").json()["signers"][0]
    assert s0["status"] == "completed" and s0["ds_ip"] == "" and s0["ds_opened_at"] is None


def test_long_details_are_trimmed(env):
    send(env)
    token, sid, submitter_id = token_of(env)
    webhook("form.completed", {"id": submitter_id, "ip": "9" * 500, "ua": "U" * 5000})
    s0 = env.get(f"/api/generated/{env.gid}/signature-request").json()["signers"][0]
    assert len(s0["ds_ip"]) == 64 and len(s0["ds_user_agent"]) == 300


# ---- the storage module on its own ----------------------------------------

def test_roundtrip_and_wrong_key(tmp_path, monkeypatch):
    monkeypatch.delenv("PHOTO_ENCRYPTION_KEY", raising=False)
    path = ps.save_image(str(tmp_path / "a"), JPEG, "secret-one")
    assert path.endswith(".enc")
    assert ps.read_image(path, "secret-one") == (JPEG, "image/jpeg")
    with pytest.raises(ps.PhotoError):
        ps.read_image(path, "secret-two")


def test_dedicated_key_overrides_the_derived_one(tmp_path, monkeypatch):
    from cryptography.fernet import Fernet
    monkeypatch.setenv("PHOTO_ENCRYPTION_KEY", Fernet.generate_key().decode())
    path = ps.save_image(str(tmp_path / "b"), PNG, "anything")
    assert ps.read_image(path, "a totally different secret") == (PNG, "image/png")


def test_sniffing():
    assert ps.sniff_mime(JPEG) == "image/jpeg"
    assert ps.sniff_mime(PNG) == "image/png"
    assert ps.sniff_mime(b"GIF89a") == ""
    assert ps.sniff_mime(b"") == ""
