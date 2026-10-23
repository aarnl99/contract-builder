"""Draft by email, end to end: only the account owner can use the drafting
address; a finished draft goes to the owner (not the client); replying with
the client's name and email is what sends the secure review link.

Run with:  venv/bin/pytest tests/
"""
import json
import os
import sys
import time

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlmodel import Session, select  # noqa: E402

from test_multi_signer import env  # noqa: E402,F401
from app import email_engine as ee  # noqa: E402
from app import main as m  # noqa: E402
from app.db import engine  # noqa: E402
from app.models import (EmailAlias, EmailDraftRequest, GeneratedContract, Placeholder, ShareLink, Template,  # noqa: E402
                        User)

PATH_SECRET = "path-secret-flow"
AUTH_OK = {"spf": "pass", "dkim": "pass", "dmarc": "pass"}


# --- parsing helpers -------------------------------------------------------

@pytest.mark.parametrize("text", [
    "Tiffany Xia, tiffany@virio.ai",
    "Send it to Tiffany Xia <tiffany@virio.ai>",
    "First name: Tiffany\nLast name: Xia\nEmail: tiffany@virio.ai",
    "Tiffany Xia\ntiffany@virio.ai\n\nThanks,\nAaron Luan",
    "Hi\n\nTiffany Xia - tiffany@virio.ai\n\nBest,\nAaron\n\nOn Mon, Oct 5 at 3:00 PM Rotely wrote:\n> old other@z.com",
])
def test_recipient_is_read_from_a_short_reply(text):
    got = ee.parse_recipient(ee.strip_quoted_reply(text), exclude=["aaron@x.test"])
    assert got == {"first_name": "Tiffany", "last_name": "Xia", "email": "tiffany@virio.ai"}


def test_incomplete_replies_leave_blanks_instead_of_guessing():
    assert ee.parse_recipient("tiffany@virio.ai")["last_name"] == ""
    assert ee.parse_recipient("Tiffany, tiffany@virio.ai")["first_name"] == ""
    assert ee.parse_recipient("Tiffany Xia")["email"] == ""
    # the owner's own address (signature, quoted text) is never taken as the client
    assert ee.parse_recipient("Tiffany Xia aaron@x.test", exclude=["aaron@x.test"])["email"] == ""


def test_only_vouched_for_senders_count_as_authenticated():
    assert ee.sender_is_authenticated({"dmarc": "pass", "dkim": "gray"})
    assert ee.sender_is_authenticated({"dkim": "pass"})
    assert not ee.sender_is_authenticated({"spf": "pass", "dkim": "fail", "dmarc": "fail"})
    assert not ee.sender_is_authenticated(None)


# --- the flow ---------------------------------------------------------------

@pytest.fixture()
def flow(env, monkeypatch):
    monkeypatch.setattr(m, "INBOUND_WEBHOOK_SECRET", PATH_SECRET)
    monkeypatch.setattr(ee, "RESEND_WEBHOOK_SECRET", "")
    mail = []
    monkeypatch.setattr(m.ee, "send_email", lambda to, subject, body, reply_to=None, attachments=None, **k:
                        mail.append({"to": to, "subject": subject, "body": body, "reply_to": reply_to, "att": attachments}))
    monkeypatch.setattr(m.ee, "extract_field_values", lambda body, placeholders: {"f0": "Acme", "f1": "12 months"})
    state = {"auth": AUTH_OK}
    monkeypatch.setattr(ee, "fetch_received_email", lambda email_id: {"text": state["text"], "authentication": state["auth"]})

    with Session(engine) as s:
        owner = s.get(User, s.get(GeneratedContract, env.gid).owner_id)
        owner_id, owner_email = owner.id, owner.email
        tpl = Template(owner_id=owner_id, name="NDA", document_type="Confidentiality", original_filename="n.docx", working_path="")
        s.add(tpl)
        s.commit()
        s.refresh(tpl)
        d = m._template_dir(owner_id, tpl.id)
        os.makedirs(d, exist_ok=True)
        doc = Document()
        doc.add_paragraph("Between {{f0}} for {{f1}}.")
        tpl.working_path = os.path.join(d, "working.docx")
        doc.save(tpl.working_path)
        s.add(tpl)
        for i in range(2):
            s.add(Placeholder(template_id=tpl.id, field_key="f%d" % i, label="Field %d" % i, order=i,
                              threshold_type="approved_list", threshold_config='{"allowed": ["a"]}',
                              preset_options_json='[{"name": "x", "text": "y"}]', required=True))
        s.commit()
    address = env.get("/api/account/email-alias").json()["address"]
    env.mail, env.state, env.owner_email, env.owner_id, env.address = mail, state, owner_email, owner_id, address
    return env


def send_in(flow, sender, subject, text, to=None, auth=None):
    flow.state["text"] = text
    if auth is not None:
        flow.state["auth"] = auth
    payload = {"type": "email.received", "data": {"email_id": "x", "from": sender, "to": [to or flow.address],
                                                   "received_for": [], "subject": subject}}
    return TestClient(m.app).post(f"/api/email/inbound/{PATH_SECRET}", json=payload)


def pending_request(flow):
    with Session(engine) as s:
        return s.exec(select(EmailDraftRequest).where(EmailDraftRequest.user_id == flow.owner_id)
                      .order_by(EmailDraftRequest.id.desc())).first()


def continuation(flow, req_id):
    with Session(engine) as s:
        a = s.exec(select(EmailAlias).where(EmailAlias.user_id == flow.owner_id)).first()
        return ee.continuation_address(a.company_slug, a.number, req_id)


def make_draft(flow):
    r = send_in(flow, flow.owner_email, "Please draft an NDA", "Client Acme, term 12 months")
    assert r.status_code == 200 and r.json() == {"ok": True}
    return pending_request(flow)


def test_finished_draft_goes_to_the_owner_only_and_invites_a_reply(flow):
    req = make_draft(flow)
    assert req.status == "awaiting_recipient" and req.generated_contract_id
    assert [x["to"] for x in flow.mail] == [flow.owner_email]
    mail = flow.mail[0]
    assert "first draft" in mail["body"].lower() and "reply to this email" in mail["body"].lower()
    assert mail["att"] and mail["att"][0]["filename"].endswith(".docx")
    assert mail["reply_to"] == continuation(flow, req.id)
    with Session(engine) as s:
        assert s.exec(select(ShareLink).where(ShareLink.generated_contract_id == req.generated_contract_id)).first() is None


def test_replying_with_the_clients_details_sends_the_secure_link(flow):
    req = make_draft(flow)
    flow.mail.clear()
    r = send_in(flow, flow.owner_email, "Re: Please draft an NDA", "Tiffany Xia, tiffany@virio.ai\n\nThanks,\nAaron",
                to=continuation(flow, req.id))
    assert r.json() == {"ok": True}
    to = {x["to"]: x for x in flow.mail}
    assert set(to) == {"tiffany@virio.ai", flow.owner_email}
    client_mail = to["tiffany@virio.ai"]
    assert "Hi Tiffany" in client_mail["body"] and "/share/" in client_mail["body"] and "Access code:" in client_mail["body"]
    assert "Tiffany Xia" in to[flow.owner_email]["body"] and "tiffany@virio.ai" in to[flow.owner_email]["body"]
    with Session(engine) as s:
        link = s.exec(select(ShareLink).where(ShareLink.generated_contract_id == req.generated_contract_id)).one()
        assert (link.client_first_name, link.client_last_name, link.client_email) == ("Tiffany", "Xia", "tiffany@virio.ai")
        assert s.get(EmailDraftRequest, req.id).status == "sent"


def test_a_second_reply_does_not_send_again(flow):
    req = make_draft(flow)
    cont = continuation(flow, req.id)
    send_in(flow, flow.owner_email, "Re: x", "Tiffany Xia, tiffany@virio.ai", to=cont)
    flow.mail.clear()
    send_in(flow, flow.owner_email, "Re: x", "Other Person, other@virio.ai", to=cont)
    assert [x["to"] for x in flow.mail] == [flow.owner_email]
    assert "already sent" in flow.mail[0]["body"].lower()
    with Session(engine) as s:
        assert len(s.exec(select(ShareLink).where(ShareLink.generated_contract_id == req.generated_contract_id)).all()) == 1


def test_missing_details_are_asked_for_and_nothing_goes_to_the_client(flow):
    req = make_draft(flow)
    flow.mail.clear()
    send_in(flow, flow.owner_email, "Re: x", "tiffany@virio.ai", to=continuation(flow, req.id))
    assert [x["to"] for x in flow.mail] == [flow.owner_email]
    assert "first name" in flow.mail[0]["body"] and "last name" in flow.mail[0]["body"]
    with Session(engine) as s:
        assert s.get(EmailDraftRequest, req.id).status == "awaiting_recipient"
        assert s.exec(select(ShareLink).where(ShareLink.generated_contract_id == req.generated_contract_id)).first() is None
    # the follow-up with everything then sends
    flow.mail.clear()
    send_in(flow, flow.owner_email, "Re: x", "Tiffany Xia, tiffany@virio.ai", to=continuation(flow, req.id))
    assert "tiffany@virio.ai" in {x["to"] for x in flow.mail}


def test_a_colleague_cannot_use_the_address_and_gets_pointed_to_sign_up(flow):
    r = send_in(flow, "colleague@acme.test", "Please draft an NDA", "Client Acme, term 12 months")
    assert r.json()["ok"] is False and "owner" in r.json()["reason"]
    assert pending_request(flow) is None
    assert [x["to"] for x in flow.mail] == ["colleague@acme.test"]
    assert "own account" in flow.mail[0]["body"] and "Acme" not in flow.mail[0]["body"]


def test_a_forged_sender_is_ignored_silently(flow):
    r = send_in(flow, "colleague@acme.test", "Please draft an NDA", "x", auth={"spf": "fail", "dkim": "fail", "dmarc": "fail"})
    assert r.json()["ok"] is False
    assert flow.mail == [] and pending_request(flow) is None


def test_a_forged_owner_address_cannot_send_anything_to_a_client(flow):
    req = make_draft(flow)
    flow.mail.clear()
    r = send_in(flow, flow.owner_email, "Re: x", "Mallory Evil, mallory@evil.test", to=continuation(flow, req.id),
                auth={"spf": "fail", "dkim": "fail", "dmarc": "fail"})
    assert r.json()["ok"] is False
    assert [x["to"] for x in flow.mail] == [flow.owner_email] and "verify" in flow.mail[0]["body"].lower()
    with Session(engine) as s:
        assert s.exec(select(ShareLink).where(ShareLink.generated_contract_id == req.generated_contract_id)).first() is None


def test_unknown_subject_reply_lists_the_available_documents(flow):
    send_in(flow, flow.owner_email, "something unrelated", "hello")
    assert "- NDA" in flow.mail[0]["body"]
